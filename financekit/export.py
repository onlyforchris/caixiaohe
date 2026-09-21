# -*- coding: utf-8 -*-
"""CSV / XLSX 导出（纯函数，无网络、无可变状态）。"""

import io

from financekit.parse import get_categories


# ---------- CSV ----------
def csv_escape(v):
    v = "" if v is None else str(v)
    if any(ch in v for ch in (",", '"', "\n")):
        v = '"' + v.replace('"', '""') + '"'
    return v


def export_csv(recs, scope):
    head = ["文件名", "所在文件夹", "金额(元)", "开票日期", "发票号码", "发票代码",
            "销售方", "费用分类", "分类依据", "状态", "重复提示", "OCR/提示", "文件路径"]
    lines = [",".join(head)]
    for r in recs:
        used = "已使用" if r["is_used"] else "待使用"
        dup = "；".join("%s[%s/%s]" % (d["basis"], d["folder"], d["fname"])
                        for d in r["dups"]) or ""
        if (scope == "used" and not r["is_used"]) or \
           (scope == "dup" and not r["dups"]) or \
           (scope == "reused" and (r["is_used"] or not any(d.get("is_used") for d in r["dups"]))) or \
           (scope == "new" and not r.get("in_watch")):
            continue
        row = [r["fname"], r["folder"], (r["amount_cents"] or 0) / 100,
               r.get("date") or "", r.get("no") or "", r.get("code") or "",
               r.get("seller") or "", r.get("cat_label") or "", r.get("cat_rule") or "",
               used, dup, r.get("warn") or "", r["path"]]
        lines.append(",".join(csv_escape(x) for x in row))
    return "\ufeff" + "\n".join(lines)


def record_in_scope(r, scope):
    if scope == "used":
        return r["is_used"]
    if scope == "new":
        return bool(r.get("in_watch"))
    if scope == "dup":
        return bool(r.get("dups"))
    if scope == "reused":
        return (not r["is_used"]) and any(d.get("is_used") for d in (r.get("dups") or []))
    return True


# ---------- XLSX ----------
export_worksheet_head = ["序号", "提交人/批次", "文件名", "发票号码", "开票日期",
                         "销售方", "金额(元)", "费用分类", "查验状态",
                         "合规校验", "退回原因", "报销状态"]


def export_worksheet(recs):
    """审核底稿：明细 + 按提交人小计 + 合计 + 退回清单，可直接打印/存档。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    rows = [r for r in recs if r.get("in_watch")]
    rows.sort(key=lambda r: (r.get("folder") or "", r.get("date") or "", r.get("fname") or ""))
    wb = Workbook()
    ws = wb.active
    ws.title = "审核底稿"
    bold = Font(bold=True)
    ws.append(export_worksheet_head)
    for c in ws[1]:
        c.font = bold
    for i, r in enumerate(rows, 1):
        issues = "；".join([*(r.get("check_issues") or []), *(r.get("ai_issues") or [])])
        ws.append([i, r.get("folder") or "", r.get("fname") or "", r.get("no") or "",
                   r.get("date") or "", r.get("seller") or "",
                   (r.get("amount_cents") or 0) / 100, r.get("cat_label") or "",
                   r.get("verify_status") or "未查验", issues,
                   r.get("reject_reason") or "",
                   "已报销" if r.get("is_used") else "待报销"])
    n = len(rows)
    start = n + 3
    ws.cell(row=start, column=1, value="按提交人/批次小计").font = bold
    sums = {}
    for r in rows:
        sums[r.get("folder") or "（未分组）"] = sums.get(r.get("folder") or "（未分组）", 0) \
            + (r.get("amount_cents") or 0)
    for j, (k, v) in enumerate(sorted(sums.items()), start + 1):
        ws.cell(row=j, column=2, value=k)
        c = ws.cell(row=j, column=7, value=v / 100)
        c.number_format = "0.00"
    total_row = start + 1 + len(sums) + 1
    ws.cell(row=total_row, column=1, value="合计").font = bold
    t = ws.cell(row=total_row, column=7, value="=SUM(G2:G%d)" % (n + 1))
    t.font = bold
    t.number_format = "0.00"
    widths = {1: 6, 2: 18, 3: 30, 4: 22, 5: 12, 6: 26, 7: 12, 8: 14, 9: 12, 10: 30, 11: 14, 12: 10}
    for k, v in widths.items():
        ws.column_dimensions[ws.cell(row=1, column=k).column_letter].width = v
    rejected = [r for r in rows if r.get("reject_reason")]
    ws2 = wb.create_sheet("退回清单")
    ws2.append(["提交人/批次", "文件名", "发票号码", "金额(元)", "退回原因"])
    for c in ws2[1]:
        c.font = bold
    for r in rejected:
        ws2.append([r.get("folder") or "", r.get("fname") or "", r.get("no") or "",
                    (r.get("amount_cents") or 0) / 100, r.get("reject_reason")])
    if not rejected:
        ws2.append(["（无退回票据）", "", "", "", ""])
    for col, w in zip("ABCDE", (18, 30, 22, 12, 16)):
        ws2.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_summary_workbook(recs):
    """分类汇总表：分类做列、金额按开票日期竖排、底部合计行，另附明细页。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    ordered = [c["label"] for c in get_categories()]
    groups = {}
    for r in recs:
        cents = r.get("amount_cents")
        if not cents:
            continue
        lbl = r.get("cat_label") or "其他/待分类"
        if lbl not in ordered:
            ordered.append(lbl)
        groups.setdefault(lbl, []).append((r.get("date") or "", cents))
    cols = [lbl for lbl in ordered if groups.get(lbl)]
    for lst in groups.values():
        lst.sort(key=lambda x: (x[0] or "9999-99-99", x[1]))

    wb = Workbook()
    ws = wb.active
    ws.title = "分类汇总"
    bold = Font(bold=True)
    n_rows = max((len(v) for v in groups.values()), default=0)
    for j, lbl in enumerate(cols, 1):
        col = get_column_letter(j)
        head = ws.cell(row=1, column=j, value=lbl)
        head.font = bold
        ws.column_dimensions[col].width = max(12, len(lbl) * 2 + 4)
        for i, (_d, cents) in enumerate(groups[lbl], 2):
            cell = ws.cell(row=i, column=j, value=cents / 100)
            cell.number_format = "0.00"
        sum_row = n_rows + 3
        total = ws.cell(row=sum_row, column=j, value="=SUM(%s2:%s%d)" % (col, col, n_rows + 1))
        total.font = bold
        total.number_format = "0.00"
    if cols:
        last = len(cols) + 1
        head = ws.cell(row=1, column=last, value="总计")
        head.font = bold
        grand = ws.cell(row=n_rows + 3, column=last,
                        value="=SUM(A%d:%s%d)" % (n_rows + 3, get_column_letter(len(cols)), n_rows + 3))
        grand.font = bold
        grand.number_format = "0.00"
        ws.column_dimensions[get_column_letter(last)].width = 14

    ws2 = wb.create_sheet("明细")
    head = ["开票日期", "金额(元)", "费用分类", "销售方", "文件名", "所在文件夹", "状态"]
    for j, h in enumerate(head, 1):
        cell = ws2.cell(row=1, column=j, value=h)
        cell.font = bold
    for i, r in enumerate(sorted(recs, key=lambda x: (x.get("date") or "9999-99-99", x.get("fname") or "")), 2):
        vals = [r.get("date") or "", (r.get("amount_cents") or 0) / 100,
                r.get("cat_label") or "", r.get("seller") or "", r.get("fname") or "",
                r.get("folder") or "", "已使用" if r["is_used"] else "待使用"]
        for j, v in enumerate(vals, 1):
            cell = ws2.cell(row=i, column=j, value=v)
            if j == 2:
                cell.number_format = "0.00"
    for j, w in enumerate((12, 11, 13, 32, 36, 20, 9), 1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    return wb


def export_xlsx(recs):
    from io import BytesIO
    buf = BytesIO()
    build_summary_workbook(recs).save(buf)
    return buf.getvalue()
