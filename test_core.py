# -*- coding: utf-8 -*-
"""核心查重与分类的最小回归检查。"""
import io
import urllib.error

import app
import update


def keys(fields, md5="same"):
    return {key for _level, key in app.fingerprint(fields, md5)}


def main():
    a = {"no": "12345678", "amount_cents": 1000}
    b = {"no": "12345678", "amount_cents": 1100}
    assert keys(a) & keys(b), "同发票号码、金额识别不一致时必须提示"
    assert "md5|same" in keys(a), "有结构化字段时也必须保留文件内容指纹"
    assert app.is_within(r"C:\used\a.pdf", r"C:\used")
    assert not app.is_within(r"C:\used_backup\a.pdf", r"C:\used")
    assert app.classify_text("项目名称：打印纸")[0] == "office"
    assert app.classify_text("项目名称：酒店住宿")[0] == "travel"
    assert app.classify_text("项目名称：停车费")[0] == "transport"
    assert app.classify_text("高德打车网约车服务")[0] == "transport"
    assert app.classify_text("员工下午茶餐饮")[0] == "benefit"
    assert app.classify_text("生产生活服务餐费")[0] == "benefit"
    assert app.classify_text("项目名称：休闲零食食品")[0] == "benefit"
    assert app.classify_text("项目名称：宴请礼品")[0] == "hospitality"
    assert app.classify_text("顺丰收派服务")[0] == "office"
    assert app.classify_text("网约车客运服务费")[0] == "transport"
    assert app.classify_text("92号车用乙醇")[0] == "transport"
    assert app.classify_text("人工智能技术服务")[0] == "service"
    assert app.classify_text("熟肉制品酱板鸭")[0] == "benefit"
    broken = "发票号码：\x002\x004\x003\x001\x002\x000\x000\x000\x000\x000\x009\x008\x007\x006\x005\x004\x003\x002\x001\x000\n2026^t09g\b07\n价税合计（小写）\x00¥\x001\x00.\x009\x006"
    fields = app.extract_fields(broken)
    assert fields["no"] == "24312000009876543210"
    assert fields["date"] == "2026-09-07"
    assert fields["amount_cents"] == 196
    # 数电票新版式：标签与值分离 + 数值逐位拆空格
    sparse = ("电子发票（普通发票） 发票号码：\n开票日期：\n购\n买\n方\n信\n息\n"
              "2 4 3 3 2 0 0 0 0 0 0 8 7 6 5 4 3 2 1 0\n"
              "2 0 2 6 年0 9 月0 9 日\n"
              "测试科技有限公司\n91330100TEST00001X\n"
              "杭州味道（杭 州）餐饮有限公司\n91330100TEST00002X\n"
              "¥5 8 . 0 0 ¥0 . 5 8\n伍拾捌圆伍角捌分 ¥5 8 . 5 8\n"
              "* 生产生活服务* 餐饮服务 1 %5 8 . 0 0 0 . 5 8")
    f2 = app.extract_fields(sparse, company_names=["测试科技有限公司"])
    assert f2["no"] == "24332000000876543210", f2["no"]
    assert f2["date"] == "2026-09-09", f2["date"]
    assert f2["amount_cents"] == 5858, f2["amount_cents"]
    assert f2["buyer"] == "测试科技有限公司"
    # 注：NFKC 归一化会把全角括号统一为半角，属于预期行为（利于查重比对一致性）
    assert f2["seller"] == "杭州味道(杭州)餐饮有限公司", f2["seller"]
    assert f2["kind"] == "数电票"
    assert not app.validate_fields(f2), app.validate_fields(f2)
    assert app.extract_fields("价税合计（大写）：壹佰圆肆角伍分 （小写）：100.45")["amount_cents"] == 10045
    party = app.extract_fields(
        "名称: 测试科技有限公司\n统一社会信用代码: 91330100TEST00001X\n"
        "名称: 杭州市西湖区测试小吃店\n统一社会信用代码: 92330100TEST00003X",
        company_names=["测试科技有限公司"])
    assert party["seller"] == "杭州市西湖区测试小吃店"
    # 个体小商户后缀（小吃店/商行等）也必须能识别
    assert "杭州测试小吃店" in app._companies(
        "91330100TEST00001X\n杭州测试小吃店\n92310100TEST00004X")
    # 杂串粘连（数电票机器可读区）：b017 不得被当成公司名前缀
    dirty = "01,32,,24317200000008021560,16.40,20260908,,b017测试科技有限公司\n" \
            "91330100TEST00001X\n杭州测试小吃店\n92310100TEST00004X"
    dcomps = app._companies(dirty)
    assert not any("b017" in c for c in dcomps), dcomps
    df = app.extract_fields(dirty, company_names=["测试科技有限公司"])
    assert df["buyer"] == "测试科技有限公司", df["buyer"]
    assert df["seller"] == "杭州测试小吃店", df["seller"]
    joined = app.extract_fields(
        "24332000007717763210\n2026年09月07日\n测试科技有限公司\n"
        "91330100TEST00001X杭州市西湖区测试烧饼店\n92330100TEST00005X",
        company_names=["测试科技有限公司"])
    assert joined["seller"] == "杭州市西湖区测试烧饼店"
    compact = "测试科技有限公司 杭州⾦测试⻝品⻔店\n91330100TEST00001X 91310000TEST00006X"
    fields = app.extract_fields(compact, company_names=["测试科技有限公司"])
    assert fields["seller"] == "杭州金测试食品门店"
    assert app.validate_categories(app.DEFAULT_CATEGORIES)
    assert update.version_key("v1.1.10") > update.version_key("1.1.9")
    assert update.protected("financekit.db-wal")
    assert update.protected("backups/ledger.json")
    assert not update.protected("financekit/paths.py")
    assert "无效" in app.friendly_ocr_error(urllib.error.HTTPError("", 401, "", {}, None))
    assert "频率" in app.friendly_ocr_error(urllib.error.HTTPError("", 429, "", {}, None))

    # xlsx 分类汇总导出
    from openpyxl import load_workbook
    recs = [
        {"path": "a", "fname": "a.pdf", "folder": "F", "amount_cents": 1470,
         "date": "2026-08-02", "cat_label": "交通费", "seller": "s1",
         "is_used": False, "dups": []},
        {"path": "b", "fname": "b.pdf", "folder": "F", "amount_cents": 200,
         "date": "2026-08-01", "cat_label": "交通费", "seller": "s2",
         "is_used": False, "dups": []},
        {"path": "c", "fname": "c.pdf", "folder": "F", "amount_cents": 77700,
         "date": "2026-08-03", "cat_label": "差旅费", "seller": "s3",
         "is_used": True, "dups": []},
        {"path": "d", "fname": "d.pdf", "folder": "F", "amount_cents": None,
         "date": None, "cat_label": "交通费", "seller": "s4",
         "is_used": False, "dups": []},
    ]
    buf = app.export_xlsx(recs)
    wb = load_workbook(filename=io.BytesIO(buf))
    ws = wb["分类汇总"]
    assert ws.cell(1, 1).value == "差旅费", "列顺序须按分类目录（差旅费在交通费前）"
    assert ws.cell(1, 2).value == "交通费"
    assert ws.cell(2, 2).value == 2, "金额按日期升序：0.8-1 的 2 元在前"
    assert ws.cell(3, 2).value == 14.7
    assert ws.cell(2, 1).value == 777
    assert ws.cell(5, 1).value == "=SUM(A2:A3)", "合计行统一覆盖数据区（空单元格按0计）"
    assert ws.cell(5, 2).value == "=SUM(B2:B3)"
    assert ws.cell(1, 3).value == "总计"
    ws2 = wb["明细"]
    assert ws2.cell(2, 4).value == "s2"
    assert app.record_in_scope({"is_used": False, "dups": [{"is_used": True}]}, "reused")
    assert not app.record_in_scope({"is_used": True, "dups": [{"is_used": True}]}, "reused")
    assert app.record_in_scope({"is_used": True, "in_watch": True, "dups": []}, "new")
    assert not app.record_in_scope({"is_used": False, "in_watch": False, "dups": []}, "new")
    assert not app.record_in_scope({"is_used": False, "dups": []}, "dup")

    # T2 字段自校验
    assert any("位数" in x for x in app.validate_fields({"no": "12345"}))
    assert not app.validate_fields({"no": "12345678"})
    assert not app.validate_fields({"no": "2" + "0" * 19})
    assert any("晚于今天" in x for x in app.validate_fields({"date": "2099-01-01"}))
    assert any("自洽" in x for x in app.validate_fields(
        {"amount_cents": 10000, "amount_excl_cents": 5000, "tax_cents": 300}))
    assert not app.validate_fields(
        {"amount_cents": 10000, "amount_excl_cents": 9434, "tax_cents": 566})

    # T5 抬头校验
    assert app.check_buyer({"buyer": "测试科技有限公司"}, ["测试科技"])[0] == "符合"
    assert app.check_buyer({"buyer": "杭州市西湖区测试小吃店"}, ["测试科技"])[0] == "不符"
    assert app.check_buyer({"buyer": None}, [])[0] == "未校验"

    # N2 永久库：无 md5 时不得生成 exact 键，否则条目互相误判
    assert not any(k.startswith("md5|") for _w, k in app.fingerprint({"no": "1" * 20}, None))
    assert any(k.startswith("md5|") for _w, k in app.fingerprint({"no": "1" * 20}, "abc"))
    assert app.used_item_key({"no": "1" * 20, "amount_cents": 100}) == "no+amt|%s|100" % ("1" * 20)
    assert app.used_item_key({"no": None, "amount_cents": None}) == ""

    # N3 历史台账导入
    hist = app.parse_history_csv("发票号码,金额,开票日期\n24312000000112345678,100.45,2026-08-01\n")
    assert len(hist) == 1 and hist[0]["no"] == "24312000000112345678" and hist[0]["amount_cents"] == 10045

    # N4 归档规范命名
    nm = app.archive_name({"date": "2026-08-01", "seller": "杭州*餐饮/店", "amount_cents": 10045,
                           "no": "24312000000112345678", "fname": "a.pdf"})
    assert nm.startswith("2026-08-01_杭州餐饮店_100.45_24312000000112345678") and nm.endswith(".pdf")

    # OFD（数电票官方格式）：zip + XML 解析，零第三方依赖
    import os
    import shutil
    import tempfile
    import zipfile
    tmpd = tempfile.mkdtemp(prefix="ofd-test-")
    try:
        ofd_p = os.path.join(tmpd, "t.ofd")
        with zipfile.ZipFile(ofd_p, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("OFD.xml",
                       '<?xml version="1.0" encoding="UTF-8"?>'
                       '<ofd:OFD xmlns:ofd="http://www.ofdspec.org/2016">'
                       '<ofd:DocBody><ofd:DocRoot>Doc_0/Document.xml</ofd:DocRoot></ofd:DocBody></ofd:OFD>')
            z.writestr("Doc_0/Document.xml",
                       '<?xml version="1.0" encoding="UTF-8"?>'
                       '<ofd:Document xmlns:ofd="http://www.ofdspec.org/2016"><ofd:Pages>'
                       '<ofd:Page ID="1" BaseLoc="Pages/Page_0/Content.xml"/></ofd:Pages></ofd:Document>')
            z.writestr("Doc_0/Pages/Page_0/Content.xml",
                       '<?xml version="1.0" encoding="UTF-8"?>'
                       '<ofd:Content xmlns:ofd="http://www.ofdspec.org/2016"><ofd:Layer ID="1">'
                       '<ofd:TextObject ID="1" Boundary="10 10 100 8"><ofd:TextCode>电子发票（普通发票）</ofd:TextCode></ofd:TextObject>'
                       '<ofd:TextObject ID="2" Boundary="10 20 100 8"><ofd:TextCode>发票号码：</ofd:TextCode></ofd:TextObject>'
                       '<ofd:TextObject ID="3" Boundary="60 20 100 8"><ofd:TextCode>24312000000112345678</ofd:TextCode></ofd:TextObject>'
                       '<ofd:TextObject ID="4" Boundary="10 30 100 8"><ofd:TextCode>开票日期：2026年09月09日</ofd:TextCode></ofd:TextObject>'
                       '<ofd:TextObject ID="5" Boundary="10 40 100 8"><ofd:TextCode>测试科技有限公司</ofd:TextCode></ofd:TextObject>'
                       '<ofd:TextObject ID="6" Boundary="10 50 100 8"><ofd:TextCode>杭州味道餐饮有限公司</ofd:TextCode></ofd:TextObject>'
                       '<ofd:TextObject ID="7" Boundary="10 60 100 8"><ofd:TextCode>价税合计（小写）¥58.58</ofd:TextCode></ofd:TextObject>'
                       '</ofd:Layer></ofd:Content>')
        otxt = app.read_ofd_text(ofd_p)
        assert "24312000000112345678" in otxt, otxt
        of = app.extract_fields(otxt, company_names=["测试科技有限公司"])
        assert of["no"] == "24312000000112345678"
        assert of["date"] == "2026-09-09"
        assert of["amount_cents"] == 5858
        assert of["kind"] == "数电票"
        assert of["seller"] == "杭州味道餐饮有限公司"
        assert ".ofd" in app.INVOICE_EXTS
        assert app.read_ofd_text(os.path.join(tmpd, "nope.ofd")) == ""
    finally:
        shutil.rmtree(tmpd, ignore_errors=True)

    # 本地识别为可选依赖：未安装时 local_ocr_ready() 必须为 False 且不抛异常
    assert isinstance(app.local_ocr_ready(), bool)

    # 「数字在前、¥ 在后」版式（滴滴/高德等电子行程发票实测）：标签被逐字拆开、
    # 金额行形如 “合 计 366.67¥ 11.00¥”“（ 小 写 ） 377.67¥叁佰…”。
    # 曾经的 bug：稀疏归一化把两笔拼成 “366.67¥11.00¥”，¥ 兜底抓出假金额 11.00。
    rev = app.extract_fields(
        "开票日期 : 2026年08月17日\n发票号码 : 24317000000998877665\n"
        "名称： 测试科技有限公司\n统一社会信用代码/纳税人识别号： 91330100TEST00001X\n"
        "名称： 杭州测试出行科技有限公司\n"
        "合 计 366.67¥ 11.00¥\n"
        "价 税 合 计 （ 大 写 ） （ 小 写 ） 377.67¥叁佰柒拾柒圆陆角柒分\n")
    assert rev["no"] == "24317000000998877665"
    assert rev["amount_cents"] == 37767, "¥ 后置版式必须以价税合计 377.67 为准，不得抓到假的 11.00"
    assert rev["amount_excl_cents"] == 36667 and rev["tax_cents"] == 1100
    assert rev["seller"] == "杭州测试出行科技有限公司"
    assert app.validate_fields(rev) == [], "377.67 = 366.67 + 11.00，自检应通过"

    # 数电票打车票：价税合计标签行与数值行被 备/注/开票人 隔开，前三条金额规则
    # 全部落空，第 4 条「合计+数字」兜底曾抓到不含税合计 568.93（v1.5.0 实测 bug）。
    # 金额必须取大写紧邻的 ¥ 小写 586.00。
    split_amt = (
        "电子发票（普通发票）\n发票号码： 24127000000400186170\n开票日期： 2026年08月17日\n"
        "购\n买\n方\n信\n息\n杭州测试出游科技有限公司\n"
        "*交通运输服务*客运服务费 568.9320388349515 1 568.93 3% 17.07\n"
        "合        计 568.93¥ 17.07¥\n"
        "价税合计（大写） （小写）\n备\n注\n开票人：\n"
        "伍佰捌拾陆圆整 ¥586.00\n测试员\n")
    fa = app.extract_fields(split_amt)
    assert fa["no"] == "24127000000400186170", fa["no"]
    assert fa["amount_cents"] == 58600, fa["amount_cents"]
    assert fa["amount_excl_cents"] == 56893 and fa["tax_cents"] == 1707
    assert app.validate_fields(fa) == [], app.validate_fields(fa)

    # 高德打车版式：¥10.96 与下一行 “10.64¥ 0.32¥” 被稀疏归一化粘成
    # “¥10.9610.64¥0.32¥”，¥ 兜底曾从中段抓出假金额 9610.64（v1.5.0 实测 bug）
    fused = (
        "电子发票（普通发票）\n发票号码：24337000000786748450\n开票日期：2026年09月07日\n"
        "价税合计（大写） （小写）\n合 计\n备\n注\n"
        "壹拾圆零玖角陆分 ¥10.96\n10.64¥ 0.32¥\n"
        "*交通运输服务*客运服务 10.640776699029126 1 10.64 3% 0.32\n费\n")
    ff = app.extract_fields(fused)
    assert ff["amount_cents"] == 1096, ff["amount_cents"]

    # 审核底稿导出
    from openpyxl import load_workbook as _lw
    wrecs = [{"path": "a", "fname": "a.pdf", "folder": "张三", "in_watch": True,
              "amount_cents": 1000, "date": "2026-08-01", "seller": "s", "no": "n1",
              "cat_label": "交通费", "verify_status": "已查验通过", "check_issues": [],
              "reject_reason": "", "is_used": False},
             {"path": "b", "fname": "b.pdf", "folder": "李四", "in_watch": True,
              "amount_cents": 2000, "date": "2026-08-02", "seller": "s", "no": "n2",
              "cat_label": "差旅费", "verify_status": "查验异常", "check_issues": ["号码位数异常"],
              "reject_reason": "疑似假票", "is_used": False}]
    wb3 = _lw(filename=io.BytesIO(app.export_worksheet(wrecs)))
    ws3 = wb3["审核底稿"]
    assert ws3.cell(1, 9).value == "查验状态"
    assert ws3.cell(2, 3).value == "a.pdf" and ws3.cell(3, 3).value == "b.pdf"
    assert wb3["退回清单"].cell(2, 3).value == "n2", "退回清单只含有退回原因的票"

    # 行程单（航空/打车行程单不是发票）：不识别、不查重、与同名发票配对展示
    assert app.is_itinerary_name("张三-1.96元-行程单.pdf")
    assert app.is_itinerary_name("电子行程报销单.jpg")
    assert not app.is_itinerary_name("高德打车电子发票.pdf")
    assert not app.is_itinerary_name("发票.pdf")
    itf = app.extract_fields("航空运输电子客票行程单 乘客 张三", "张三-1.96元-行程单.pdf")
    assert itf["itinerary"] and itf["kind"] == "行程单"
    assert itf["no"] is None and itf["amount_cents"] is None and itf["seller"] is None
    itf2 = app.extract_fields("航空运输电子客票行程单\n合计 560.00", "download.pdf")
    assert itf2["itinerary"] and itf2["amount_cents"] is None, itf2
    inv_txt = "电子发票（普通发票） 发票号码：24312000000012345680 备注：行程单见附件"
    assert not app.extract_fields(inv_txt, "发票.pdf").get("itinerary")
    assert app._stem_core("【优e出租-17.85元-1个行程】高德打车电子发票.pdf") == \
        app._stem_core("【优e出租-17.85元-1个行程】高德打车行程单.pdf")
    inv_r = {"path": "p1", "fname": "【优e出租-17.85元-1个行程】高德打车电子发票.pdf",
             "folder": "202610"}
    it_r = {"path": "p2", "fname": "【优e出租-17.85元-1个行程】高德打车行程单.pdf",
            "folder": "202610", "itinerary": True}
    app.pair_itineraries([inv_r, it_r])
    assert it_r["pair"]["path"] == "p1"
    assert inv_r["itins"] == [{"path": "p2", "fname": it_r["fname"]}]
    solo = {"path": "p3", "fname": "北京-上海行程单.pdf", "folder": "x", "itinerary": True}
    app.pair_itineraries([inv_r, solo])
    assert "pair" not in solo
    assert app.fingerprint_for(app._itinerary_fields(), "abc") == []
    pf, pf_warn = app.parse_file(r"D:\不存在目录\行程单零成本.pdf")
    assert pf["itinerary"] and pf["kind"] == "行程单" and not pf_warn

    # 台账幽灵清理 + 移动后内容继承（v1.5.2）：存储路径全部指向临时目录
    import json
    from financekit import invoice as _inv
    from financekit import store as _store
    old = (_inv.CONFIG, _store.LEDGER, _store.DB_PATH, _store.USED_LEDGER,
           _store.BACKUP_DIR, _store._MIGRATED, _inv._RAPID_FAILED)
    tmpd2 = tempfile.mkdtemp(prefix="scan-test-")
    try:
        _inv._RAPID_FAILED = True  # 离线：跳过本地识别引擎初始化
        _inv.CONFIG = os.path.join(tmpd2, "config.json")
        json.dump({"company_names": []}, open(_inv.CONFIG, "w", encoding="utf-8"))
        _store.LEDGER = os.path.join(tmpd2, "ledger.json")
        _store.DB_PATH = os.path.join(tmpd2, "ledger.db")
        _store.USED_LEDGER = os.path.join(tmpd2, "used.json")
        _store.BACKUP_DIR = os.path.join(tmpd2, "backups")
        _store._MIGRATED = False

        watch1 = os.path.join(tmpd2, "202609")
        watch2 = os.path.join(tmpd2, "202610")
        os.makedirs(watch1)
        os.makedirs(watch2)

        def ghost(name, **kw):
            p = os.path.join(watch1, name)
            r = {"path": p, "folder": "202609", "fname": name, "no": None,
                 "amount_cents": None, "itinerary": False, "fp": [], "md5": None,
                 "verify_status": "未查验", "used_override": False,
                 "used_override_set": False, "closed": False, "file_missing": False}
            r.update(kw)
            return p, r

        g1, r1g = ghost("ghost1.pdf")  # 未报销幽灵：必须被清理
        g2, r2g = ghost("ghost2.pdf", used_override=True, used_override_set=True)
        g3, r3g = ghost("ghost3.pdf", closed=True)
        _inv.save_ledger({"version": _inv.ENGINE_VER,
                          "records": {g1: r1g, g2: r2g, g3: r3g}})

        def make_ofd(path):
            with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
                z.writestr("OFD.xml",
                           '<?xml version="1.0" encoding="UTF-8"?>'
                           '<ofd:OFD xmlns:ofd="http://www.ofdspec.org/2016">'
                           '<ofd:DocBody><ofd:DocRoot>Doc_0/Document.xml</ofd:DocRoot></ofd:DocBody></ofd:OFD>')
                z.writestr("Doc_0/Document.xml",
                           '<?xml version="1.0" encoding="UTF-8"?>'
                           '<ofd:Document xmlns:ofd="http://www.ofdspec.org/2016"><ofd:Pages>'
                           '<ofd:Page ID="1" BaseLoc="Pages/Page_0/Content.xml"/></ofd:Pages></ofd:Document>')
                z.writestr("Doc_0/Pages/Page_0/Content.xml",
                           '<?xml version="1.0" encoding="UTF-8"?>'
                           '<ofd:Content xmlns:ofd="http://www.ofdspec.org/2016"><ofd:Layer ID="1">'
                           '<ofd:TextObject ID="1" Boundary="10 10 100 8"><ofd:TextCode>电子发票（普通发票）</ofd:TextCode></ofd:TextObject>'
                           '<ofd:TextObject ID="2" Boundary="10 20 100 8"><ofd:TextCode>发票号码：</ofd:TextCode></ofd:TextObject>'
                           '<ofd:TextObject ID="3" Boundary="60 20 100 8"><ofd:TextCode>24312000000112345678</ofd:TextCode></ofd:TextObject>'
                           '<ofd:TextObject ID="4" Boundary="10 30 100 8"><ofd:TextCode>开票日期：2026年09月09日</ofd:TextCode></ofd:TextObject>'
                           '<ofd:TextObject ID="5" Boundary="10 50 100 8"><ofd:TextCode>杭州味道餐饮有限公司</ofd:TextCode></ofd:TextObject>'
                           '<ofd:TextObject ID="6" Boundary="10 60 100 8"><ofd:TextCode>价税合计（小写）¥58.58</ofd:TextCode></ofd:TextObject>'
                           '</ofd:Layer></ofd:Content>')

        ofd1 = os.path.join(watch1, "inv1.ofd")
        make_ofd(ofd1)
        led, _u1, _c1 = app.build_scan([watch1], [], False, "glm-4v-flash")
        p1 = os.path.normpath(ofd1)
        assert p1 in led["records"], "新文件必须入库"
        assert g1 not in led["records"], \
            "未报销且未结案的幽灵记录必须在扫描时清理（即使 verify_status 有值）"
        assert led["records"][g2]["file_missing"] is True, "已标记使用的幽灵必须保留"
        assert led["records"][g3]["file_missing"] is True, "已结案的幽灵必须保留"

        # 整月归档场景：文件原样移动到新月份目录后，同 md5 新路径直接继承识别结果
        r1 = led["records"][p1]
        ofd2 = os.path.join(watch2, "inv1.ofd")
        shutil.move(ofd1, ofd2)
        led2, _u2, _c2 = app.build_scan([watch2], [watch1], False, "glm-4v-flash")
        p2 = os.path.normpath(ofd2)
        assert p2 in led2["records"], "移动后的文件必须出现在台账里"
        r2 = led2["records"][p2]
        assert r2["no"] == "24312000000112345678" and r2["amount_cents"] == 5858, \
            "继承记录的字段必须与原识别结果一致"
        assert r2["first_seen"] == r1["first_seen"], \
            "first_seen 不变才证明是继承而非重新识别"
        assert p1 not in led2["records"], "旧路径的幽灵记录应被清理"
        assert r2["dups"] == [], "继承后不得与任何残留记录误判重复"
    finally:
        (_inv.CONFIG, _store.LEDGER, _store.DB_PATH, _store.USED_LEDGER,
         _store.BACKUP_DIR, _store._MIGRATED, _inv._RAPID_FAILED) = old
        shutil.rmtree(tmpd2, ignore_errors=True)

    # v1.5.3 数电票稀疏版式（标签聚在顶部、值在下方）：销售方以“馆/铺/中心”
    # 结尾时必须能识别。
    def _sparse(no, date, bname, btax, sname, stax):
        return ("电子发票（普通发票） 发票号码：\n开票日期：\n购\n买\n方\n信\n"
                "息 统一社会信用代码/纳税人识别号：\n销\n售\n方\n信\n"
                "息 统一社会信用代码/纳税人识别号：\n名称： 名称：\n"
                "项目名称 规格型号 单 位 数 量 单 价 金 额 税率/征收率 税 额\n"
                "合 计\n价税合计（大写） （小写）\n备\n注\n开票人：\n"
                + no + "\n" + date + "\n" + bname + "\n" + btax + "\n"
                + sname + "\n" + stax + "\n¥51.49 ¥0.51\n伍拾贰圆整 ¥52.00\n")

    f1 = app.extract_fields(
        _sparse("24312000001845948181", "2026年09月16日",
                "测试科技有限公司", "91330100TEST00001X",
                "杭州西湖测试泡馍馆", "92330100TEST00007X"),
        "dzfp.pdf", company_names=["测试科技有限公司"])
    assert f1["buyer"] == "测试科技有限公司", f1["buyer"]
    assert f1["seller"] == "杭州西湖测试泡馍馆", f1["seller"]
    # 买方不是本公司时按出现顺序取第一个为买方
    f2 = app.extract_fields(
        _sparse("24334000000253374946", "2026年06月28日",
                "杭州测试社会工作发展中心", "52330100TEST00008X",
                "杭州市西湖区测试牛肉面馆", "92330100TEST00009X"),
        "manual.pdf", company_names=["测试科技有限公司"])
    assert f2["buyer"] == "杭州测试社会工作发展中心", f2["buyer"]
    assert f2["seller"] == "杭州市西湖区测试牛肉面馆", f2["seller"]
    # “铺”后缀
    f3 = app.extract_fields(
        _sparse("24334000000367914796", "2026年09月11日",
                "测试科技有限公司", "91330100TEST00001X",
                "杭州钱塘测试粥铺", "92330100TEST00010X"),
        "manual.pdf", company_names=["测试科技有限公司"])
    assert f3["seller"] == "杭州钱塘测试粥铺", f3["seller"]

    print("core checks: ok")


class _ModelDown(Exception):
    """模拟某档模型返回 5xx：只看 .code，用于驱动降级。"""
    code = 503


def _test_vision_model_fallback():
    """视觉模型自动选择与降级：第一档失败应选中第二档并锁定到本进程。"""
    import financekit.invoice as _inv  # 下划线开头的符号 import * 不会带进来

    real_key = _inv.zhipu_key
    real_chat = _inv.zhipu_chat
    real_active = _inv._ACTIVE_VISION_MODEL
    calls = []

    def fake_key():
        return "test.fake.key"

    def fake_chat(messages, model=None, timeout=60, max_tokens=None):
        calls.append(model)
        if model == "glm-4.6v-flash":
            raise _ModelDown("glm-4.6v-flash 暂时不可用")
        return "OK"

    # 直接 patch 源模块：pick_vision_model → _probe_model → zhipu_chat 全在同一模块内做全局查找
    _inv.zhipu_key = fake_key
    _inv.zhipu_chat = fake_chat
    _inv._ACTIVE_VISION_MODEL = None
    try:
        # 第一次调用：应探测链上两档，跳过第一档后锁定第二档
        m1 = _inv.pick_vision_model()
        assert m1 == "glm-4v-flash", f"首档失败时应回退到第二档，实际选中 {m1!r}"
        assert calls == ["glm-4.6v-flash", "glm-4v-flash"], \
            f"探测顺序不对：{calls}"

        # 锁定后再次请求：不应再发探测请求
        calls.clear()
        m2 = _inv.pick_vision_model()
        assert m2 == "glm-4v-flash", f"锁定后必须返回已选模型，实际 {m2!r}"
        assert calls == [], f"锁定后不应再探测，实际仍调用：{calls}"
    finally:
        _inv.zhipu_key = real_key
        _inv.zhipu_chat = real_chat
        _inv._ACTIVE_VISION_MODEL = real_active


if __name__ == "__main__":
    _test_vision_model_fallback()
    main()
